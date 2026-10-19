#!/usr/bin/env python3
"""Auditoría previa a una entrega LP-R03 22.9 kV.

No modifica M01 ni publica ingeniería. Su resultado es explícito: BLOQUEADA si
hay documentos de 10 kV o si faltan archivos nuevos debidamente regenerados.
Se ejecuta con Python con ezdxf, openpyxl y PyMuPDF instalados.
"""
from pathlib import Path
import argparse
import json
import re

PROJECT = Path(__file__).resolve().parents[2]

def auditar(project, directorio):
    checks = []
    def add(id, ok, detail, severity="ERROR"):
        checks.append({"id":id,"ok":bool(ok),"detail":str(detail),"severity":severity})
    try:
        import ezdxf
        from openpyxl import load_workbook
        import fitz
    except ImportError as exc:
        add("DEPENDENCIAS",False,exc)
        return checks
    cfgpath = project / "datos/lp-r03-229kv/escenario.json"
    cfg = json.loads(cfgpath.read_text(encoding="utf-8")) if cfgpath.exists() else {}
    add("TENSION_ADOPTADA", cfg.get("tension_nominal_linea_kv")==22.9,
        f'Valor declarado: {cfg.get("tension_nominal_linea_kv")}')
    add("ESTADO_HONESTO", cfg.get("aprobado_para_construccion") is False,
        "No indicar obra aprobada sin inspecciones y ensayos.")
    # No confundir CAD M01 congelado con el diseño nuevo.
    cad = directorio / "04_Planos/Linea_Primaria_R03_229KV.dxf"
    if not cad.exists():
        add("CAD_R03",False,"Falta DWG/DXF técnico nuevo, no basta A01 de estilo.")
    else:
        try:
            d=ezdxf.readfile(cad)
            ids=[x.get_attrib_text("ID") for x in d.modelspace().query("INSERT") if x.get_attrib_text("ID")]
            add("111_APOYOS",len(ids)==111 and len(set(ids))==111,f"Apoyos: {len(ids)}, únicos: {len(set(ids))}")
            lays=[l.name for l in d.layouts if l.name.startswith("LP-")]
            add("PLANOS_LP",len(lays)>=19,f"Layouts: {len(lays)}")
            add("AUDIT_CAD",not d.audit().errors,"Auditoría DXF sin errores" if not d.audit().errors else "CAD corrupto")
            labels=" ".join(x.dxf.text for layout in d.layouts for x in layout.query("TEXT") if hasattr(x.dxf,"text"))
            # Históricos R02 pueden existir en documentación, pero NO con carácter de lámina vigente.
            add("ROTULOS_R03", bool(re.search(r"22[.,]9\s*kV",labels,re.I)),
                "Se requiere tensión nominal 22.9 kV en las láminas nuevas.")
            add("ROTULOS_ANTIGUOS", not bool(re.search(r"LP-R02\s*\|\s*18/10/2026",labels)),
                "No dejar los rótulos R02 como si la lámina fuera R03.")
        except Exception as exc:
            add("CAD_ABRIR",False,exc)
    xls=directorio/"03_Metrados/Metrados_R03_229KV.xlsx"
    if xls.exists():
        try:
            wb=load_workbook(xls,read_only=True,data_only=True)
            add("METRADOS_POSTES",any("Inventario" in s for s in wb.sheetnames),
                "Se requiere inventario por poste de M01 y adaptación verificada.")
            add("FICHA_POSTE","Ficha_poste" in wb.sheetnames,
                "Conservar ficha interactiva del poste M01.")
            add("MATRIZ_ARMADOS","Matriz_armados" in wb.sheetnames,
                "Conservar despiece por armado M01.")
            wb.close()
        except Exception as exc: add("XLSX_ABRIR",False,exc)
    else:
        add("METRADOS_R03",False,"No existe XLSX M01 recalculado con armados de 22.9 kV.")
    pdf=directorio/"00_Expediente/Expediente_R03_229KV.pdf"
    if pdf.exists():
        try:
            doc=fitz.open(pdf)
            txt=" ".join(p.get_text() for p in doc)
            add("PDF_PAGINAS",len(doc)>70,f"{len(doc)} páginas; expediente requiere anexos y planos.")
            add("PDF_22_9KV",bool(re.search(r"22[.,]9\s*kV",txt,re.I)),"Tensión nominal en documento.")
            add("PDF_TENSION_VIGENTE",not bool(re.search(r"(tensión de cálculo|tensión nominal|caso base)\s*:?\s*10\s*kV",txt,re.I)),
                "Rechazar expresiones que presenten 10kV como caso nominal vigente.")
            doc.close()
        except Exception as exc: add("PDF_ABRIR",False,exc)
    else:
        add("EXPEDIENTE_R03",False,"No se ha regenerado expediente integrado R03.")
    for name in ("02_Calculos/Calculos_R03_229KV.xlsx",
                 "06_Presupuesto/Presupuesto_R03_229KV.xlsx",
                 "04_Planos/Planos_R03_229KV.pdf"):
        add("EXISTE_"+name.split("/")[-1],(directorio/name).exists(),name)
    return checks

def main():
    a=argparse.ArgumentParser()
    a.add_argument("--entrega",type=Path,required=True)
    a.add_argument("--informe",type=Path)
    x=a.parse_args()
    checks=auditar(PROJECT,x.entrega)
    failed=[c for c in checks if not c["ok"] and c["severity"]=="ERROR"]
    result={"revision":"LP-R03-229KV","resultado":"BLOQUEADA" if failed else "CONTROLES AUTOMATICOS OK — REVISION HUMANA PENDIENTE",
            "fallos":len(failed),"controles":checks,"aprobado_para_construccion":False}
    if x.informe:
        if x.informe.exists(): raise FileExistsError("No sobrescribir informe anterior.")
        x.informe.parent.mkdir(parents=True,exist_ok=True)
        x.informe.write_text(json.dumps(result,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    print(json.dumps(result,ensure_ascii=False,indent=2))
    raise SystemExit(1 if failed else 0)
if __name__=="__main__":main()
