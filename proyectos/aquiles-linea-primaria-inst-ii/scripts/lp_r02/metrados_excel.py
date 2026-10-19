"""Excel M01 con catálogo editable, inventario completo y ficha seleccionable."""
from pathlib import Path
import shutil
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation
from metrados_componentes import cargar, medir
from rutas_entrega import ruta_activa

AZUL='254D73';CELESTE='EAF2F8';BLANCO='FFFFFF';ENTRADA='225DA8'
GRUPOS=['Postes','Crucetas y soportes','Aisladores','Pernos y accesorios',
        'Conexiones y amarres','Retenida','Puesta a tierra de apoyo','Accesorios incluidos']


def orden_materiales(cat):
    return sorted(cat['materiales'],key=lambda k:(GRUPOS.index(cat['materiales'][k]['grupo']),list(cat['materiales']).index(k)))


def hoja(w,nombre,titulo,cabeceras,anchos):
    if nombre in w:del w[nombre]
    s=w.create_sheet(nombre);s.sheet_view.showGridLines=False
    s.cell(2,1,titulo).font=Font(name='Arial',size=15,bold=True,color=AZUL)
    for j,(h,ancho) in enumerate(zip(cabeceras,anchos),1):
        x=s.cell(7,j,h);x.font=Font(name='Arial',size=10,bold=True,color=BLANCO)
        x.fill=PatternFill('solid',fgColor=AZUL);x.alignment=Alignment(wrap_text=True,vertical='center')
        s.column_dimensions[x.column_letter].width=ancho
    s.row_dimensions[7].height=32;s.freeze_panes='D8'
    s.page_setup.orientation='landscape';s.page_setup.paperSize=s.PAPERSIZE_A3
    s.page_setup.fitToWidth=1;s.page_setup.fitToHeight=0;s.sheet_properties.pageSetUpPr.fitToPage=True
    s.print_title_rows='1:7';s.oddFooter.center.text='Limbani–Huancasayani · M01 · &P de &N'
    return s


def estilo(s,fila,valores):
    for j,v in enumerate(valores,1):
        x=s.cell(fila,j,v);x.font=Font(name='Arial',size=9,color='20252B')
        x.alignment=Alignment(vertical='top',wrap_text=j in (2,3,5,10,11,12))
        if isinstance(v,(float,int)) or (isinstance(v,str) and v.startswith('=')):
            x.number_format='#,##0.000;-#,##0.000;"—"'
        if fila%2==0:x.fill=PatternFill('solid',fgColor='F3F7FB')


def actualizar_basicos(w,criterios):
    """Separa puentes eléctricos de los ramales mecánicos en los tres libros."""
    if 'Criterios_M01' in w:del w['Criterios_M01']
    cr=w.create_sheet('Criterios_M01')
    cr.append(['Criterio M01','Unidad','Valor','Estado'])
    cr.append(['Ángulo ideal de retenida','°',criterios['ang_retenida_grados'],'Modelo ideal LP-R02'])
    cr.append(['Colas de cable por ramal','m',criterios['colas_retenida_por_ramal_m'],'Hipótesis académica explícita'])
    cr.append(['Conexión enterrada y colas PAT','m',criterios['cola_bajante_pat_m'],'Hipótesis de metrado, sin Wenner'])
    cr.append(['Paso nominal de fijación PAT','m',criterios['paso_fijacion_pat_m'],'Hipótesis de montaje para metrado'])
    cr.column_dimensions['A'].width=42;cr.column_dimensions['D'].width=48
    s=w['Postes'];s['N6']='Armado M01';s['O6']='Puentes/colas (ud)'
    s['P6']='Cable RI/ramal (m)';s['Q6']='Cu PAT/apoyo (m)';s['R6']='Fijaciones PAT (jgo)';s['S6']='Tendido asignado (m)'
    for col in 'NOPQRS':s.column_dimensions[col].width=19
    for r in range(7,118):
        s[f'N{r}']=f'=IF(I{r}="SUSPENSION","CAC-S",IF(I{r}="ANGULO","CAC-A",IF(I{r}="TERMINAL","CAC-T",IF(OR(I{r}="ANCLAJE",I{r}="ANCLAJE_PROPUESTO"),"CAC-R","SIN-RECETA"))))'
        s[f'O{r}']=f'=IF(OR(I{r}="ANCLAJE",I{r}="ANCLAJE_PROPUESTO",I{r}="TERMINAL"),Parametros!$E$5,0)'
        s[f'M{r}']=f'=O{r}*Parametros!$E$6'
        s[f'P{r}']=f'=G{r}/SIN(RADIANS(Criterios_M01!$C$2))+Criterios_M01!$C$3'
        s[f'Q{r}']=f'=G{r}+Criterios_M01!$C$4'
        s[f'R{r}']=f'=ROUNDUP(G{r}/Criterios_M01!$C$5,0)+1'
        s[f'S{r}']=f'=(SUMIF(Vanos!$B$7:$B$116,A{r},Vanos!$N$7:$N$116)+SUMIF(Vanos!$C$7:$C$116,A{r},Vanos!$N$7:$N$116))/2'
    if 'Resumen' in w and w['Resumen']['B25'].value=='Puentes y colas':
        w['Resumen']['D25']='3 fases × (38 retenciones + 2 terminales) × 2 m'
        w['Resumen']['B31']='Puentes por interrupción del conductor; retenidas se miden por separado.'
    w.calculation.fullCalcOnLoad=True;w.calculation.forceFullCalc=True


def ampliar(w,base):
    cat,c=cargar();m=medir(base,cat,c);keys=orden_materiales(cat)
    materiales=cat['materiales'];armados=cat['armados'];codigos=['CAC-S','CAC-A','CAC-R','CAC-T','RI-A','PAT-P']
    mat=hoja(w,'Matriz_armados','CONSUMO POR ARMADO · COEFICIENTES UNITARIOS',
             ['Material','Descripción','Un.']+codigos+['Procedencia / alcance'],[16,66,8]+[12]*6+[65])
    mat['A4']='Azul: coeficientes editables. La unidad corresponde a una configuración o a un ramal RI-A.'
    mapa={k:i+8 for i,k in enumerate(keys)}
    for k in keys:
        e=materiales[k];r=mapa[k]
        vals=[k,e['descripcion'],e['unidad']]+[armados[a]['coef'].get(k,0) for a in codigos]
        vals += [e['codigo_referencia']+'; '+('ACCESORIO INCLUIDO, no compra independiente' if e['grupo']=='Accesorios incluidos' else 'Ver bases y adaptación CAC en Trazabilidad')]
        estilo(mat,r,vals);mat.row_dimensions[r].height=33
        formato='#,##0.000;-#,##0.000;"—"' if e['unidad']=='m' else '#,##0;-#,##0;"—"'
        for j in range(4,10):
            mat.cell(r,j).number_format=formato
            mat.cell(r,j).font=Font(name='Arial',size=10,color=ENTRADA)
            mat.cell(r,j).fill=PatternFill('solid',fgColor='EAF3FF')
            mat.cell(r,j).comment=Comment('Coeficiente de '+codigos[j-4]+'. '+armados[codigos[j-4]]['fuente_cantidades']+'. Las adaptaciones son hipótesis de metrado M01.','Proyecto')
    mat.auto_filter.ref=f'A7:J{7+len(keys)}'

    inv=hoja(w,'Inventario','INVENTARIO DESARROLLADO POR POSTE Y COMPONENTE',
             ['Poste','Material','Descripción','Un.','Código ref.','Armado','Retenidas','PAT','Poste','Total','Nivel','Armado asignado'],
             [11,16,66,8,24,13,13,13,13,14,14,20])
    inv['A4']='Filtra Poste y Total > 0. COMPRA: cantidades de suministro; INCLUIDO: accesorios ya contenidos en pernos/abrazaderas.'
    inv['A5']='CONTROL: distribución contable del conductor; su compra total figura en Resumen y Resumen_materiales.'
    filas={};r=8
    for p in base['postes']:
        rr=7+next(i for i,q in enumerate(base['postes']) if q['id']==p['id'])
        for k in keys+['AAAC-T','AAAC-P']:
            filas[(p['id'],k)]=r;r+=1
    fin=r-1
    for i,p in enumerate(base['postes']):
        pp=7+i
        for k in keys+['AAAC-T','AAAC-P']:
            r=filas[(p['id'],k)]
            if k.startswith('AAAC-'):
                val=f'=Postes!{("S" if k=="AAAC-T" else "M")}{pp}'
                estilo(inv,r,[p['id'],k,'Conductor '+('tendido (asignación de vanos)' if k=='AAAC-T' else 'puentes o colas'),'m','AAAC 50 mm²',0,0,0,0,val,'CONTROL',f'=Postes!N{pp}'])
                continue
            e=materiales[k];q=mapa[k];nivel='INCLUIDO' if e['grupo']=='Accesorios incluidos' else 'COMPRA'
            fa=f'=INDEX(Matriz_armados!$D${q}:$G${q},1,MATCH(Postes!N{pp},Matriz_armados!$D$7:$G$7,0))*Postes!L{pp}'
            fr=f'=Matriz_armados!H{q}*Postes!J{pp}'
            ft=f'=Matriz_armados!I{q}'
            fp=f'=IF(Postes!E{pp}={k[1:]},1,0)' if k in ['P13','P15','P17'] else '=0'
            if k=='GUY':fr=f'=Postes!P{pp}*Postes!J{pp}'
            if k=='CU-PAT':ft=f'=Postes!Q{pp}'
            if k=='FI-PAT':ft=f'=Postes!R{pp}'
            total=f'=SUM(F{r}:I{r})'
            if nivel=='INCLUIDO':
                terms=[f'J{filas[(p["id"],parent)]}*{v["incluye"][k]}'
                       for parent,v in materiales.items() if k in v['incluye']]
                fa=fr=ft=fp='=0';total='='+'+'.join(terms) if terms else '=0'
            estilo(inv,r,[p['id'],k,e['descripcion'],e['unidad'],e['codigo_referencia'],fa,fr,ft,fp,total,nivel,f'=Postes!N{pp}'])
            for j in range(6,11):inv.cell(r,j).number_format='#,##0.000;-#,##0.000;"—"' if e['unidad']=='m' else '#,##0;-#,##0;"—"'
            inv.row_dimensions[r].height=27
            if nivel=='INCLUIDO':
                for celda in inv[r]:celda.font=Font(name='Arial',size=9,color='65727D',italic=True)
    inv.auto_filter.ref=f'A7:L{fin}'
    J=f'Inventario!$J$8:$J${fin}';B=f'Inventario!$B$8:$B${fin}';K=f'Inventario!$K$8:$K${fin}';A=f'Inventario!$A$8:$A${fin}'

    resumen=hoja(w,'Resumen_materiales','RESUMEN GENERAL DE MATERIALES PARA SUMINISTRO',
                 ['Código','Material','Un.','Cantidad de compra','Referencia / criterio'],[16,75,8,22,70])
    incluidos=hoja(w,'Accesorios_incluidos','ACCESORIOS YA INCLUIDOS CON PERNOS Y ABRAZADERAS',
                   ['Código','Accesorio','Un.','Cantidad incluida','Tratamiento'],[16,70,8,22,65])
    resumen['A4']='Los precios por conjunto de armado/retenida ya incluyen su despiece. No volver a valorizar cada pieza encima del conjunto.'
    incluidos['A4']='Control físico: estas cantidades no se vuelven a comprar ni valorizar por separado.'
    a=b=8
    for k in keys:
        e=materiales[k];inc=e['grupo']=='Accesorios incluidos';s=incluidos if inc else resumen;n=b if inc else a
        nivel='INCLUIDO' if inc else 'COMPRA'
        estilo(s,n,[k,e['descripcion'],e['unidad'],f'=SUMIFS({J},{B},A{n},{K},"{nivel}")',
                     'Incluido en conjunto; no sumar como compra adicional' if inc else e['codigo_referencia']])
        s.cell(n,4).number_format='#,##0.000;-#,##0.000;"—"' if e['unidad']=='m' else '#,##0;-#,##0;"—"'
        s.row_dimensions[n].height=30
        if inc:b+=1
        else:a+=1
    estilo(resumen,a,['AAAC-50','Conductor AAAC 50 mm² para suministro','m','=Resumen!E27','Tendido + puentes/colas + reserva + redondeo'])
    resumen.auto_filter.ref=f'A7:E{a}';incluidos.auto_filter.ref=f'A7:E{b-1}'

    ficha=hoja(w,'Ficha_poste','FICHA DE INVENTARIO · SELECCIONA EL POSTE EN B4',
               ['Código','Componente','Un.','Armado','Retenidas','PAT','Poste','Total','Nivel'],
               [16,70,8,13,13,13,13,14,14])
    ficha['A4']='Poste:';ficha['B4']=base['postes'][0]['id'];ficha['B4'].fill=PatternFill('solid',fgColor='DCEBFF')
    ficha['C4']='Armado:';ficha['D4']="=INDEX(Postes!$N$7:$N$117,MATCH(B4,Postes!$A$7:$A$117,0))"
    ficha['A5']='Las filas con guión tienen cantidad cero; los accesorios INCLUIDO se muestran sin compra adicional.'
    v=DataValidation(type='list',formula1='INDIRECT("\'Postes\'!$A$7:$A$117")');v.error='Selecciona un identificador válido.';v.showErrorMessage=True
    ficha.add_data_validation(v);v.add(ficha['B4']);ficha.freeze_panes='D8'
    for r,k in enumerate(keys+['AAAC-T','AAAC-P'],8):
        e=materiales.get(k,dict(descripcion='Conductor '+('tendido asignado' if k=='AAAC-T' else 'puentes/colas'),unidad='m',grupo='CONTROL'))
        formula=[]
        for col in 'FGHIJ':
            rango=f'Inventario!${col}$8:${col}${fin}'
            formula.append(f'=SUMIFS({rango},{A},$B$4,{B},A{r})')
        nivel='CONTROL' if k.startswith('AAAC-') else ('INCLUIDO' if e['grupo']=='Accesorios incluidos' else 'COMPRA')
        estilo(ficha,r,[k,e['descripcion'],e['unidad']]+formula+[nivel]);ficha.row_dimensions[r].height=29
        for j in range(4,9):ficha.cell(r,j).number_format='#,##0.000;-#,##0.000;"—"' if e['unidad']=='m' else '#,##0;-#,##0;"—"'
    ficha.auto_filter.ref=f'A7:I{7+len(keys)+2}'

    traza=hoja(w,'Trazabilidad','BASES DE MEDICIÓN Y PROCEDENCIA', ['N.º','Criterio / adaptación'],[8,150])
    notas=cat['adaptaciones']+[
        'Puentes/colas: tres fases por 38 retenciones y dos terminales; longitud adoptada 2 m por fase. Suspensión/ángulo de paso no interrumpen el conductor.',
        'Conductor tendido: mitad de cada vano asignada a cada poste extremo sólo para conciliación. No prescribe cortes ni empalmes.',
        'Perno completo: incluye los accesorios indicados por su receta. La hoja Accesorios_incluidos es un control analítico, no otra lista de compra.',
        'Para un cambio de función/altura edita Postes; para un consumo unitario edita Matriz_armados. Inventario, Ficha_poste y los resúmenes dependen de esas entradas.',
        'Las fichas PDF representan la versión publicada. Regenerar el expediente si se cambia una entrada del Excel.',
        'El presupuesto conserva precios referenciales por conjunto; el detalle nuevo no constituye una cotización comercial de cada perno.'
    ]
    notas += [f['titulo']+'. '+f['url']+' · SHA256 '+f['sha256'] for f in cat['fuentes']]
    for i,t in enumerate(notas,8):estilo(traza,i,[i-7,t]);traza.row_dimensions[i].height=46
    for i,f in enumerate(cat['fuentes'],8+len(notas)-len(cat['fuentes'])):
        traza.cell(i,2).hyperlink=f['url'];traza.cell(i,2).font=Font(name='Arial',size=9,color=ENTRADA,underline='single')
    guia=hoja(w,'Guia','METRADOS M01 · RUTA DE LECTURA', ['Paso','Contenido'],[12,120])
    for i,t in enumerate([
        'Resumen: postes, armados, retenidas y conductor total.',
        'Ficha_poste: selecciona un identificador en B4 y consulta sus componentes.',
        'Inventario: filtra por poste, componente y cantidad mayor que cero.',
        'Matriz_armados: consumo unitario por configuración, con códigos de la referencia.',
        'Resumen_materiales: compra global, conciliada con los inventarios.',
        'Accesorios_incluidos: tuercas, contratuercas y arandelas ya incluidas con los conjuntos.',
        'Postes y Vanos: ubicación, función y desarrollo de las longitudes.',
        'Trazabilidad: hipótesis de adaptación CAC, límites del catálogo y fuentes.'
    ],8):estilo(guia,i,[i-7,t]);guia.row_dimensions[i].height=36
    w._sheets=[w['Guia'],w['Resumen'],w['Ficha_poste'],w['Resumen_materiales'],w['Inventario'],w['Matriz_armados'],w['Accesorios_incluidos'],w['Postes'],w['Vanos'],w['Parametros'],w['Criterios_M01'],w['Trazabilidad']]
    w.active=0
    return m


def preparar_libros(proyecto, salida, base):
    _,c=cargar();salida=Path(salida)
    for nombre,carpeta in [('Calculos','01_Calculos'),('Metrados','02_Metrados'),('Presupuesto','05_Presupuesto')]:
        destino=salida/carpeta/(nombre+'.xlsx');destino.parent.mkdir(parents=True,exist_ok=True)
        shutil.copy2(ruta_activa(Path(carpeta)/(nombre+'.xlsx')),destino)
        w=load_workbook(destino);actualizar_basicos(w,c)
        if nombre=='Metrados':ampliar(w,base)
        w.save(destino);w.close()
    return salida
