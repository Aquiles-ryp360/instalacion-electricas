"""Expediente único: redacción técnica, ecuaciones LaTeX y metrados desarrollados."""
from pathlib import Path
import csv, hashlib, json, math, shutil, sys
import openpyxl,pymupdf
from pypdf import PdfReader,PdfWriter

P=Path(__file__).resolve().parents[2];ROOT=P.parents[1]
sys.path.insert(0,str(ROOT))
from herramientas.documentos.formato_aquiles import compilar,tex
from herramientas.lineas_primarias.revision import flujo_sistema
from metrados_componentes import actualizar_resultados_para_documento, capitulo as capitulo_metrado

B=ROOT/'build'/P.name/'presentacion-final';O=B/'Proyecto_Linea_Primaria';T=B/'latex'
R=json.loads((P/'revisiones/lp-r02/03-calculos/resultados.json').read_text(encoding='utf-8'))
C=json.loads((P/'datos/lp-r02/parametros.json').read_text(encoding='utf-8'))
R=actualizar_resultados_para_documento(R)
F=next(x for x in R['comparacion_tensiones'] if x['kV']==10 and x['toma_pu']==1.05)


def n(x,d=3):return f'{x:,.{d}f}'.replace(',',r'\,')


def tabla(head,rows,weights=None,size='small'):
    weights=weights or [1]*len(head);sw=sum(weights)
    spec=''.join('L{\\dimexpr '+str(round(w/sw,6))+r'\linewidth-2\tabcolsep\relax}' for w in weights)
    h=' & '.join(r'\textbf{'+tex(x)+'}' for x in head)+r' \\ \midrule'
    lines=['\\begingroup\\'+size,r'\begin{longtable}{'+spec+'}',r'\toprule\rowcolor{modelblue!12}',h,r'\endfirsthead',r'\toprule\rowcolor{modelblue!12}',h,r'\endhead']
    lines.extend(' & '.join(str(x) for x in row)+r' \\' for row in rows)
    return '\n'.join(lines+[r'\bottomrule\end{longtable}\endgroup',''])


def preparar():
    T.mkdir(exist_ok=True);(T/'capitulos').mkdir(exist_ok=True);(T/'plantilla').mkdir(exist_ok=True);(T/'figuras').mkdir(exist_ok=True)
    for name in ['estilo.tex','caratula.tex']:shutil.copy2(ROOT/'recursos/plantillas/expediente-aquiles'/name,T/'plantilla'/name)
    shutil.copy2(ROOT/'recursos/identidad/escudo-unap.png',T/'figuras/escudo-unap.png')
    if not (B/'ubicacion/Plano_Ubicacion.pdf').is_file():
        from ubicacion_memoria import generar
        generar(B)
    shutil.copy2(B/'ubicacion/Plano_Ubicacion.pdf',T/'figuras/ubicacion.pdf')
    macros=dict(ProyectoCabecera='PROYECTO DE LÍNEA PRIMARIA LIMBANI - HUANCASAYANI, SANDIA - PUNO',
        CursoCabecera=r'\shortstack{INSTALACIONES\\ELÉCTRICAS II}',DocumentoPie='LÍNEA PRIMARIA LIMBANI - HUANCASAYANI',AutorDocumento='Aquiles Taylor Ramos Yapo',
        CursoCaratula='INSTALACIONES ELÉCTRICAS II',TituloCaratula='PROYECTO: ``LÍNEA PRIMARIA LIMBANI - HUANCASAYANI\'\'',
        UbicacionCaratula='DISTRITO DE LIMBANI - PROVINCIA DE SANDIA - DEPARTAMENTO DE PUNO',TipoDocumentoCaratula='EXPEDIENTE DEL PROYECTO ACADÉMICO')
    raw=''.join('\\newcommand{\\'+k+'}{'+(v if k=='CursoCabecera' else tex(v))+'}\n' for k,v in macros.items())
    raw+=r'''\newcommand{\FichaCaratula}{
\textbf{Estudiante:} & Aquiles Taylor Ramos Yapo\\
\textbf{Código:} & 227368\\
\textbf{Curso:} & Instalaciones Eléctricas II\\
\textbf{Ubicación:} & Limbani - Huancasayani, Sandia, Puno\\
\textbf{Sistema:} & Aislado, trifásico, 60 Hz\\
\textbf{Documento:} & Memoria, cálculos, metrados y planos\\
\textbf{Fecha:} & 19 de octubre de 2026\\}
'''
    (T/'metadatos.tex').write_text(raw,encoding='utf-8')


def contenido():
    partes={}
    partes['01-memoria']=r'''\chapter{MEMORIA DESCRIPTIVA}
\section{Objeto del proyecto}
El proyecto desarrolla una línea primaria trifásica para transportar energía desde la transformación de salida de la minicentral de Limbani hasta la subestación de llegada en Huancasayani, provincia de Sandia, Puno. El sistema opera de forma aislada. El diseño de las redes secundarias y acometidas de usuarios queda fuera de este expediente.

El alcance académico comprende el corredor georreferenciado, la demanda de transporte, el conductor, la transformación, el comportamiento eléctrico y mecánico, las cantidades de suministro y montaje, el presupuesto referencial y las láminas de planta, perfil y esquemas. Las cantidades y resultados de las hojas de cálculo se incorporan en los capítulos correspondientes.

\section{Configuración del sistema}
\begin{equation}
\text{MCH}\;\longrightarrow\;T_1\;\longrightarrow\;\text{LP}\;\longrightarrow\;T_2\;\longrightarrow\;\text{Huancasayani}
\end{equation}
La línea se representa por un eje en planta, correspondiente a una terna de tres conductores AAAC de 50 mm$^2$. El caso de compatibilidad se calcula a 10 kV y se contrasta con 20, 22 y 22.9 kV. Los transformadores de salida y llegada se modelan con 250 kVA netos disponibles en el emplazamiento. La relación física y las prestaciones reales de los equipos no se infieren a partir del caso numérico.

\section{Ficha del emplazamiento}
'''+tabla(['Parámetro','Definición'],[
 ['Departamento / provincia / distrito','Puno / Sandia / Limbani'],['Sistema de referencia','WGS 84 / UTM zona 19 Sur, EPSG:32719'],
 ['Inicio del corredor',r'$E=425\,715.000\ \mathrm{m}$; $N=8\,436\,650.000\ \mathrm{m}$'],
 ['Final del corredor',r'$E='+n(R['postes'][-1]['e'])+r'\ \mathrm{m}$; $N='+n(R['postes'][-1]['n'])+r'\ \mathrm{m}$'],
 ['Longitud horizontal',n(R['longitud_horizontal_m'])+' m'],['Número de apoyos / vanos','111 / 110'],['Modo de operación','Sistema aislado, sin interconexión inicial con el SEIN']], [1.2,2.8])+r'''
\section{Procedencia y condiciones del diseño}
Las coordenadas y el perfil proceden del corredor de gabinete, cartografía pública, catastro aportado y elevaciones DEM de aproximadamente 30 m. No constituyen un levantamiento topográfico de campo. Los emplazamientos de T1 y T2 son posiciones académicas del corredor; no se atribuye titularidad ni servidumbre aprobada a esos puntos.

Las cargas de usuarios son proyecciones y asignaciones de diseño. Los datos de conductor, transformadores, generador y capacidad transversal de postes son datos tipo. La verificación resistente presentada cubre la flexión idealizada; no acredita capacidad axial, interacción N--M, geotecnia ni rigidez y pretensión reales de las retenidas. El anclaje R02-P019-2 presenta interferencia en la cartografía de referencia. Esta condición y las cargas ascendentes se identifican en las láminas y en la interpretación mecánica.

La potencia nominal de referencia de 200 kW se utiliza para evaluar balances, sin convertirla en potencia firme de estiaje. El expediente corresponde al diseño académico de la línea primaria y conserva independencia respecto del dimensionamiento del curso de Centrales Hidroeléctricas.
'''
    partes['02-demanda']=r'''\chapter{DEMANDA Y BALANCE DE POTENCIA}
\section{Proyección de usuarios}
La base corresponde a las viviendas ocupadas del Directorio INEI 2017, tomo 4, páginas 1713--1714, para los tres registros de Huancasayani empleados en el ámbito de estudio. Se aplica crecimiento geométrico de 1\% anual hasta 2026, con redondeo superior por localidad:
\begin{equation}
N_i=\left\lceil V_{i,2017}(1+r)^{2026-2017}\right\rceil,
\qquad r=0.01.
\end{equation}
La vivienda proyectada representa un usuario potencial del escenario. No se suma la población distrital de Limbani a los registros de Huancasayani.
'''+tabla(['Localidad','Viviendas 2017','Usuarios 2026','kW por usuario','Simult.','Demanda kW'],[[tex(x['localidad']),n(x['viviendas_base'],0),n(x['usuarios_potenciales'],0),n(x['calificacion_kW'],1),n(x['simultaneidad'],2),n(x['demanda_kW'],3)] for x in R['demanda']],[2.4,1,1,1,1,1.2])+r'''
\section{Demanda doméstica y complementaria}
\begin{align}
P_{\mathrm{dom}}&=\sum_i N_i\,p_i\,f_s\\
&=104(0.6)(0.8)+154(0.6)(0.8)+95(0.4)(0.8)\notag\\
&=154.240\ \mathrm{kW}.\notag
\end{align}
\begin{equation}
P_{\mathrm{AP}}=60(0.0541)=3.246\ \mathrm{kW}.
\end{equation}
Los servicios se asignan en 3 kW para salud, 3 kW para educación y comunicaciones, 3 kW para agua/bombeo, 4 kW para comercio/talleres y 1 kW para servicios comunales. Son cargas académicas, con suma de 14 kW.
\begin{equation}
\boxed{P_D=154.240+3.246+14.000=171.486\ \mathrm{kW}}
\end{equation}
\begin{align}
Q_D&=P_D\tan(\arccos 0.9)=83.054\ \mathrm{kvar},\\
S_D&=\frac{P_D}{0.9}=190.540\ \mathrm{kVA}.
\end{align}
\section{Balance con la fuente aislada}
La potencia de generación requerida incluye la pérdida en la línea, ambos transformadores y 5 kW de auxiliares. Para el caso desarrollado se obtiene:
\begin{equation}
P_{G,\mathrm{req}}=P_D+\Delta P_{T1}+\Delta P_{\mathrm{LP}}+\Delta P_{T2}+2P_0+P_{\mathrm{aux}}
'''+r'='+n(F['P_generacion_kW'])+r'\ \mathrm{kW}.'+r'''
\end{equation}
En esta expresión $\Delta P_{T1}$ y $\Delta P_{T2}$ representan exclusivamente pérdidas de cobre. Los consumos en vacío $2P_0$ se contabilizan una sola vez en los nodos del modelo.
'''+tabla(['Oferta nominal kW','Reserva Limbani kW','Necesidad total kW','Margen kW'],[[n(x['oferta_hipotetica_kW'],0),n(x['reserva_limbani_kW'],0),n(x['necesidad_kW']),n(x['margen_kW'])] for x in R['balance_fuente'] if x['kV']==10])+r'''
Con 200 kW y sin reserva simultánea para Limbani, el margen nominal es 16.674 kW. Al reservar 25 kW para Limbani, el balance arroja un déficit de 8.326 kW. Por ello la operación compartida se rige por el balance de potencia disponible y la gestión de carga; el valor nominal de 200 kW no garantiza el abastecimiento permanente de ambas demandas.
\begin{equation}
P_{G,\mathrm{req}}(P_H)+P_{\mathrm{Limbani}}\leq P_{\mathrm{disponible}}.
\end{equation}
\section{Crecimiento de demanda}
Se aplica $P_{D,t}=P_{D,2026}(1.01)^{t-2026}$ al perfil de estudio. La transformación se contrasta con la potencia aparente de fuente, además de la potencia activa.
'''+tabla(['Año','Carga kW','Generación kW','Fuente kVA','BT V'],[[str(x['anio']),n(x['P_carga_kW']),n(x['P_generador_kW']),n(x['S_generador_kVA']),n(x['BT_V'])] for x in R['proyeccion'] if x['kV']==10])
    partes['03-electrico']=r'''\chapter{CÁLCULOS ELÉCTRICOS}
\section{Conductor y longitud eléctrica}
Se adopta conductor AAAC 6201-T81 de sección nominal 50 mm$^2$, sección real de cálculo 49.5 mm$^2$, diámetro 9 mm y masa 135 kg/km. La longitud por fase se obtiene de la catenaria EDS de cada vano; no se reemplaza por la proyección horizontal.
\begin{align}
L_{\mathrm{XY}}&=\sum_{i=1}^{110}L_i='''+n(R['longitud_horizontal_m'])+r'''\ \mathrm{m},\\
\ell&=\sum_{i=1}^{110}L_{c,i}='''+n(R['longitud_electrica_m'])+r'''\ \mathrm{m}.
\end{align}
\section{Resistencia y reactancia}
\begin{align}
R_{75}&=R_{20}[1+\alpha_R(75-20)]\frac{\ell}{1000}\\
&=0.671[1+0.0036(55)](6.758120)\notag\\
&=\boxed{5.432569\ \Omega/\mathrm{fase}}.\notag
\end{align}
Con separaciones horizontales de 1.2, 1.2 y 2.4 m:
\begin{align}
D_m&=(D_{ab}D_{bc}D_{ca})^{1/3}=1.511905\ \mathrm{m},\\
\mathrm{GMR}&=0.7788\frac{0.009}{2}=0.0035046\ \mathrm{m},\\
X&=2\pi f(2\times10^{-4})\ln\!\left(\frac{D_m}{\mathrm{GMR}}\right)\frac{\ell}{1000}\\
&=\boxed{3.091467\ \Omega/\mathrm{fase}}.\notag
\end{align}
\section{Modelo de transformación}
Cada transformador se representa por impedancia serie y consumo en vacío. Se emplean $S_n=250$ kVA, $u_k=0.06$ pu, $P_{cu,n}=3.1$ kW, $P_0=0.45$ kW y $Q_0=2.5$ kvar.
\begin{align}
r_T&=\frac{P_{cu,n}}{S_n}=0.0124\ \mathrm{pu},\\
x_T&=\sqrt{u_k^2-r_T^2}=0.058705\ \mathrm{pu},\\
Z_b&=\frac{V_n^2}{S_n}=400\ \Omega,\\
Z_T&=(r_T+jx_T)Z_b=4.960+j23.482\ \Omega.
\end{align}
Las pérdidas y reactancias son datos tipo del modelo. La elevación efectiva de fuente $t=1.05$ pu representa una condición equivalente de regulación, sin atribuirle una posición física de toma no documentada.
\section{Flujo de potencia con T1, línea y T2}
El modelo trifásico equilibrado se resuelve por fase mediante barrido atrás/adelante. La tensión de fuente se fija en $V_0=10\,000(1.05)/\sqrt{3}$ V y las cargas se representan como potencias PQ:
\begin{align}
S_1&=S_2=450+j2500\ \mathrm{VA},\\
S_3&=171\,486+j83\,054.460\ \mathrm{VA},\\
I_j^{(k)}&=\sum_{m=j}^{3}\left(\frac{S_m}{3V_m^{(k)}}\right)^{*},\\
\widehat V_{j+1}^{(k+1)}&=\widehat V_j^{(k+1)}-Z_jI_j^{(k)},\\
V_j^{(k+1)}&=0.7\widehat V_j^{(k+1)}+0.3V_j^{(k)}.
\end{align}
La iteración se detiene con diferencia de tensión inferior a $10^{-8}$ V. El libro de cálculo desarrolla 60 iteraciones visibles, con partes real e imaginaria, sin referencias circulares.
'''+tabla(['Magnitud','Resultado'],[
 ['MT de salida de T1',n(F['MT_salida_kV'])+' kV'],['MT de llegada a T2',n(F['MT_llegada_kV'])+' kV'],['BT de llegada',n(F['BT_llegada_V'])+' V'],
 ['Corriente en la línea',n(F['corrientes_A'][1],6)+' A'],['Pérdida activa de línea',n(F['perdida_linea_kW'],6)+' kW'],['Generación requerida',n(F['P_generacion_kW'],6)+' kW'],['Potencia aparente de fuente',n(F['S_generador_kVA'],6)+' kVA']], [2,1])+r'''
\begin{align}
\Delta P_{\mathrm{LP}}&=3|I_2|^2R=2.101854\ \mathrm{kW},\\
\Delta V_{\mathrm{LP}}&=\frac{|V_1|-|V_2|}{|V_1|}\,100=1.207337\%,\\
V_{\mathrm{BT}}&=\sqrt{3}|V_3|\frac{380}{10\,000}=371.066630\ \mathrm{V}.
\end{align}
La caída anterior corresponde sólo a la línea. No debe confundirse con la regulación total generador--T1--línea--T2. Con fuente nominal equivalente sin elevación, la BT calculada es 350.291 V; la diferencia justifica analizar la regulación del conjunto.
\section{Comparación de tensiones}
'''+tabla(['Tensión nominal kV','Fuente pu','BT llegada V','Pérdida LP kW','Fuente kW'],[[n(x['kV'],1),n(x['toma_pu'],3),n(x['BT_llegada_V']),n(x['perdida_linea_kW']),n(x['P_generacion_kW'])] for x in R['comparacion_tensiones']])+r'''
El aumento de tensión reduce la corriente y las pérdidas de la línea. No elimina el déficit de potencia activa de una fuente aislada ni demuestra compatibilidad de equipos existentes. La alternativa 22.9 kV se conserva como contraste; el caso de cálculo de este expediente permanece identificado en 10 kV.
'''
    partes['04-protecciones']=r'''\chapter{CORTOCIRCUITO Y PROTECCIONES}
\section{Corrientes trifásicas de referencia}
Se representa el generador de 250 kVA por $R_a=0.02$ pu y tres reactancias: $X_d''=0.20$ pu, $X_d'=0.30$ pu y $X_d=1.80$ pu. Se calculan las impedancias de secuencia positiva referidas a MT:
\begin{align}
Z_G&=(R_a+jX_d)\frac{V_n^2}{S_G},\\
I_{k3,\mathrm{MT}}&=\frac{V_n}{\sqrt{3}|Z_G+Z_{T1}+Z_{\mathrm{LP}}|},\\
I_{k3,\mathrm{BT}}&=\frac{V_n}{\sqrt{3}|Z_G+Z_{T1}+Z_{\mathrm{LP}}+Z_{T2}|}\frac{V_n}{V_{\mathrm{BT},n}}.
\end{align}
Para el punto MT de salida se omite la impedancia de línea. La tabla distingue el instante de cortocircuito y el punto eléctrico; no utiliza un único valor de corriente para toda la instalación.
'''+tabla(['Estado','MT salida A','MT llegada A','BT llegada A'],[[tex(estado.title()),* [n(next(x['Ik3_A'] for x in R['cortocircuito'] if x['kV']==10 and x['estado']==estado and x['punto']==punto)) for punto in ['MT_SALIDA','MT_LLEGADA','BT_LLEGADA']]] for estado in ['SUBTRANSITORIO','TRANSITORIO','SOSTENIDO']])+r'''
\section{Esquema funcional de protección}
La corriente sostenida de falla MT, aproximadamente 7.731 A en llegada, es inferior a la corriente de carga del escenario. Por ello la sobrecorriente convencional no se considera suficiente por sí sola para la fuente débil. El esquema funcional incorpora protección dependiente de tensión (51V), supervisión de tensión (27/59), frecuencia (81), falla a tierra y protección térmica, con disparo de interruptores y enclavamientos del sistema aislado.

Se prevén TC/TT compatibles con la sensibilidad requerida, pararrayos en los extremos y transformación, seccionamiento visible y enclavamiento de puesta a tierra para mantenimiento. La relación de TC/TT y las curvas de relé no se deducen de corrientes tipo ni se presentan como ajustes de un equipo instalado.
\section{Puesta a tierra y neutro}
El caso de neutro resistivo se expresa con $I_N=5$ A y $t_N=10$ s como base académica. La resistencia equivalente a 10 kV sería:
\begin{align}
R_N&=\frac{10\,000}{\sqrt{3}(5)}=1\,154.701\ \Omega,\\
P_N&=I_N^2R_N=28.868\ \mathrm{kW},\\
W_N&=P_Nt_N=288.675\ \mathrm{kJ}.
\end{align}
El uso de transformador de neutro o conexión zigzag depende del grupo de transformación real. El valor resistivo no dimensiona por sí solo el aislamiento, régimen térmico ni tensiones de paso y contacto. El presupuesto incluye los conjuntos de puesta a tierra como unidades de montaje, sin atribuir resistividad medida al terreno.
'''
    partes['05-mecanica']=r'''\chapter{CÁLCULOS MECÁNICOS}
\section{Estados de carga y datos de conductor}
La carga de rotura tipo es RTS = 14\,700 N, el módulo elástico $E=62$ GPa, la sección real $A=49.5$ mm$^2$ y el coeficiente térmico $\alpha=2.3\times10^{-5}\ {}^\circ\mathrm{C}^{-1}$. El estado inicial se define a 15 °C:
\begin{align}
H_0&=0.15(14\,700)=2\,205\ \mathrm{N},\\
q_0&=\frac{135}{1000}(9.80665)=1.323898\ \mathrm{N/m},\\
EA&=62\times10^9(49.5\times10^{-6})=3\,069\,000\ \mathrm{N}.
\end{align}
'''+tabla(['Estado','T °C','Viento m/s','Hielo mm','Fluencia'],[[tex(x['id']),n(x['temperatura_c'],0),n(x['viento_ms'],1),n(x['hielo_mm'],0),n(x['creep'],6)] for x in C['clima']])+r'''
El peso adicional de hielo y la presión de viento se incorporan por estado:
\begin{align}
d_h&=d+2e_h,\\
q_v&=q_0+\rho_hg\frac{\pi}{4}(d_h^2-d^2),\qquad \rho_h=913\ \mathrm{kg/m^3},\\
p_v&=0.455v^2,\qquad q_w=p_vd_h.
\end{align}
La presión anterior es la formulación de estudio usada por el modelo. El viento de 29 m/s y el hielo se asignan a los estados de carga, mientras que el frío de --5 °C y la fluencia final de 345 microdeformaciones son sensibilidades académicas. No son registros medidos en el corredor.
\section{Catenaria con desnivel y viento}
En el plano de carga resultante, con $L$ horizontal, $\Delta z$ entre amarres y $q=\sqrt{q_v^2+q_w^2}$:
\begin{align}
d_y&=\Delta z\frac{q_v}{q},\qquad
d_x=\sqrt{L^2+\Delta z^2-d_y^2},\\
a&=\frac{H}{q},\qquad
u=\operatorname{asinh}\!\left[\frac{d_y}{2a\sinh(d_x/2a)}\right],\\
x_0&=\frac{d_x}{2}-au,\qquad
u_0=-\frac{x_0}{a},\quad u_1=\frac{d_x-x_0}{a},\\
L_c&=a[\sinh(u_1)-\sinh(u_0)],\\
T_0&=H\cosh(u_0),\qquad T_1=H\cosh(u_1).
\end{align}
El cambio de estado conserva la longitud natural del conductor bajo la aproximación elástica de primer orden:
\begin{align}
L_n&=L_c-\frac{1}{EA}\int T\,ds,\\
\int T\,ds&=H\left\{\frac{d_x}{2}+\frac{a}{4}[\sinh(2u_1)-\sinh(2u_0)]\right\},\\
L_{n,2}&=L_{n,1}[1+\alpha(T_2-T_1)+\varepsilon_c].
\end{align}
La tracción de cada estado se resuelve por bisección. Para aisladores rígidos, cada vano conserva su propia condición geométrica; no se impone una tracción horizontal común a vanos desiguales.
\section{Ejemplo desarrollado: vano V001}
Para P001--P002, $L=47.146597$ m y $\Delta z=14.315387$ m. En EDS se toma $q=q_0$ y $H=H_0$:
\begin{equation}
a=\frac{2205}{1.32389775}=1665.536\ \mathrm{m},\qquad L_c=49.273525\ \mathrm{m}.
\end{equation}
La longitud de suministro de la terna de ese vano es $3L_c=147.820576$ m. La diferencia frente a $3L$ procede del desnivel y de la curvatura del conductor.
'''+tabla(['Estado','H N','T0 N','T1 N','Flecha m','Separación m'],[[tex(x['estado']),n(x['H_N'],1),n(x['T_izq_N'],1),n(x['T_der_N'],1),n(x['flecha_vertical_m']),n(x['separacion_perfil_heredado_m'])] for x in R['vanos'][0]['estados'].values()])+r'''
\section{Momentos y capacidad tipo de postes}
Se combinan fuerzas de extremo de los vanos adyacentes con su signo y dirección. El modelo aplica factores de estudio $\gamma_T=1.65$, $\gamma_V=1.5$ y $\gamma_W=2.5$. Se incluye el viento sobre el poste cónico. La capacidad tipo de flexión se expresa:
\begin{equation}
M_{Rd}=\phi\,k_r\,(10W)\,(h_p-e_p-e_0)/1000.
\end{equation}
Con $W=600$ daN, $\phi=0.65$, $k_r=2$, $h_p=13$ m, $e_p=1.9$ m y $e_0=0.15$ m:
\begin{equation}
M_{Rd}=0.65(2)(6000)(10.95)/1000=85.410\ \mathrm{kN\,m}.
\end{equation}
'''+tabla(['Poste m','Empotr. m','Brazo m','Capacidad tipo kN·m'],[[n(h,0),n(e,1),n(h-e-.15,2),n(.65*2*6000*(h-e-.15)/1000)] for h,e in [(13,1.9),(15,2.1),(17,2.1)]])+r'''
\begin{equation}
\eta_M=\frac{M_{Ed}}{M_{Rd}}.
\end{equation}
El máximo cociente de flexión del esquema con retenidas ideales es 0.598. Este resultado corresponde a la envolvente calculada y la capacidad transversal tipo; no es una verificación de carga axial ni de interacción N--M. Se mantienen 46 apoyos con reacción vertical ascendente en algún estado y 16 retenciones adicionales en el esquema académico.
\section{Retenidas y anclajes}
Las componentes horizontales longitudinales se equilibran mediante ramales ideales. Para ángulo vertical $\theta=45^\circ$:
\begin{align}
T_R&=\frac{F_H}{\cos\theta},\\
T_{\mathrm{rot,adop}}&=10\left\lceil\frac{\max(60,T_R/0.9)}{10}\right\rceil\ \mathrm{kN},\\
\eta_R&=\frac{T_R}{0.9T_{\mathrm{rot,adop}}}.
\end{align}
La resistencia mínima de suministro se conserva en 60 kN. La coordenada geométrica de cada anclaje no equivale a capacidad geotécnica; el modelo no introduce rigidez ni pretensión de cables de retenida.
\section{Flecha y separación al terreno}
El mínimo calculado respecto al perfil DEM heredado es 7.549 m, superior al margen académico de comparación de 7.5 m. La comparación se limita al perfil disponible y no acredita un gálibo topográfico real, la envolvente lateral por viento ni la ausencia de obstáculos. Las láminas muestran las cotas de gabinete y las condiciones mecánicas que sustentan esta interpretación.
'''
    partes['06-metrados']=capitulo_metrado(R,tabla,tex,n)
    partes['07-presupuesto']=presupuesto_tex()
    partes['08-especificaciones']=r'''\chapter{ESPECIFICACIONES Y EJECUCIÓN}
\section{Equipos y materiales}
El conductor es AAAC 6201-T81 de 50 mm$^2$ nominales. Los postes se agrupan en 13, 15 y 17 m con capacidad transversal tipo de 600 daN para la verificación presentada. Las retenidas se especifican con resistencia mínima de rotura de 60 kN. Aisladores, crucetas, grapas y pernos corresponden a la función de paso, ángulo, retención o terminal indicada por apoyo.

Los transformadores se modelan con 250 kVA netos en sitio, considerando la altitud como condición de prestación. La maniobra y protección se organizan en los extremos de línea y transformación. La tensión de aislamiento y relación física de equipos dependen de la alternativa nominal; los cálculos no equiparan una relación de 22 kV con otra de 22.9 kV.
\section{Montaje y control de calidad}
La ejecución física comprende replanteo, liberación de accesos y servidumbres, cimentaciones, izado, instalación de armados, tendido por poleas, regulación de flechas, continuidad eléctrica, puesta a tierra y pruebas de maniobra/protección. La tabla de catenaria es una base de cálculo a temperaturas y cargas definidas; no sustituye la tabla de tendido asociada al material efectivamente suministrado.

Los controles comprenden certificados de conductor y herrajes, identificación de postes, integridad de aisladores, torque de conexiones, continuidad de tierras, relación y polaridad de TC/TT, aislamiento, relación de transformación y pruebas funcionales de disparo e interbloqueos. El sistema aislado exige comprobar la respuesta de frecuencia y tensión frente a escalones de carga y pérdida de generación.
\section{Referencias normativas y documentales}
Se distingue el marco de instalaciones de suministro de las instalaciones interiores. El CNE Suministro sustenta los criterios de línea y la interfaz de equipos; el CNE Utilización corresponde a instalaciones interiores y servicios auxiliares dentro de su ámbito. El RNE y las disposiciones de seguridad y servidumbres se consideran según la naturaleza de la obra.

Las referencias documentales utilizadas en el desarrollo son:
\begin{enumerate}
\item MINEM, Código Nacional de Electricidad -- Suministro 2011, RM 214-2011-MEM/DM. Copia pública consultada en la base de diseño.
\item MINEM, RD 018-2003-EM/DGE, Bases para el diseño de líneas y redes primarias para electrificación rural.
\item MINEM, RD 024-2003-EM/DGE, Especificaciones técnicas de soportes normalizados para líneas y redes primarias. Las láminas adjuntas son referencias documentales vectorizadas.
\item Electro Puno, Plan Anual de Contrataciones 2024, p.59: familia de crucetas metálicas para postes CAC; alcance de uso declarado en el metrado M01.
\item Instituto Geográfico Nacional, servicio de límites referenciales departamentales, provinciales y distritales, consultado el 19/10/2026: base de ubicación UB-01.
\item INEI, Directorio Nacional de Centros Poblados 2017, tomo 4, registros del ámbito de Huancasayani.
\item Datos de cartografía OSM, catastro aportado y DEM público empleados en el corredor de gabinete.
\end{enumerate}
La referencia de armados no acredita automáticamente una adaptación a poste CAC ni una lista de fabricación. Las hojas conservan su título y lista de componentes, y se separan de las láminas propias de la línea.
'''
    partes['09-conclusiones']=r'''\chapter{RESULTADOS DEL DISEÑO}
La línea primaria académica se desarrolla sobre un corredor horizontal de 6.696 km, con 111 apoyos y 110 vanos. La terna AAAC 50 mm$^2$ demanda '''+n(R['conductor']['tendido_terna_m'])+r''' m de tendido, '''+n(R['conductor']['puentes_provision_m'],0)+r''' m de puentes/colas y '''+n(R['conductor']['reserva_m'],0)+r''' m de reserva; la cantidad de suministro redondeada es '''+n(R['conductor']['total_compra_estudio_m'],0)+r''' m. Los apoyos se distribuyen en 72 de 13 m, 30 de 15 m y 9 de 17 m, con 111 armados y 160 ramales de retenida. El metrado M01 incorpora el despiece por componente y sus fichas individuales.

La demanda máxima del escenario es 171.486 kW y 190.540 kVA. Con el caso de 10 kV y fuente equivalente a 1.05 pu, la corriente de línea es 11.356 A, la pérdida activa es 2.102 kW y la BT de llegada es 371.067 V. El balance requiere 183.326 kW en generación y 212.481 kVA de fuente. Los resultados incorporan ambos transformadores y auxiliares.

La oferta nominal de 200 kW deja 16.674 kW de margen sin reserva simultánea para Limbani. Con reserva de 25 kW, el déficit es 8.326 kW; la operación aislada debe limitar la carga atendida a la potencia efectivamente disponible. Las alternativas de mayor tensión reducen pérdidas, sin resolver por sí solas la limitación de generación.

El esquema mecánico arroja un cociente máximo de flexión tipo de 0.598 e identifica 46 apoyos con carga ascendente. Se incorporan 16 retenciones adicionales al esquema de estudio. La separación mínima de 7.549 m se refiere al perfil DEM empleado, con las condiciones de interpretación definidas en la memoria.

La entrega reúne en un único expediente la memoria, las ecuaciones, los metrados, el presupuesto referencial y los anexos gráficos. Los libros de cálculo desarrollan las cantidades y verificaciones con fórmulas editables y mantienen sus resúmenes en PDF. Los resultados corresponden a las bases académicas explícitas del documento.
'''
    titulos=['Planta general del corredor','Unifilar y delimitación eléctrica','Tensiones y pérdidas: comparación','Protección, neutro y puesta a tierra']
    titulos += [f'Planta y perfil P{a:03}--P{b:03}' for a,b in [(1,17),(17,33),(33,49),(49,65),(65,81),(81,97),(97,111)]]
    titulos += [f'Planilla de apoyos {a}--{b}' for a,b in [(1,24),(25,48),(49,72),(73,96),(97,111)]]
    titulos += ['Armados: funciones y esfuerzos','Esquemas de cimentación y puesta a tierra','Cruces cartográficos']
    catalogo=json.loads((P/'entregables/linea-primaria-r02-presentacion/09-armados-referencia/catalogo-cad.json').read_text(encoding='utf-8'))
    partes['10-indice-grafico']=r'''\chapter*{ÍNDICE DE PLANOS Y ARMADOS}
\addcontentsline{toc}{chapter}{ÍNDICE DE PLANOS Y ARMADOS}
\section*{Anexo A. Láminas de la línea primaria}
Las láminas se incorporan inmediatamente después de la memoria. El código impreso identifica la misma presentación en el CAD editable.
'''+tabla(['Código','Lámina'],[['UB-01','Ubicación política y geográfica: Perú--Puno--Sandia--Limbani']]+[[f'LP-{i+1:02}',tex(t)] for i,t in enumerate(titulos)],[1,4])+r'''
\section*{Anexo B. Referencias documentales DGE}
Las hojas siguientes conservan las vistas y listas de componentes de su referencia. No constituyen una adaptación constructiva aprobada al corredor ni piezas paramétricas de fabricación.
'''+tabla(['Lámina DGE','Descripción'],[[tex(x['lamina']),tex(x['titulo'])] for x in catalogo],[1,4],'footnotesize')
    from memoria_descriptiva import ampliar
    partes['01-memoria']=ampliar(partes['01-memoria'],R,C,tabla,tex,n)
    for key,value in partes.items():(T/'capitulos'/(key+'.tex')).write_text(value,encoding='utf-8')
    main=r'''\documentclass[12pt,letterpaper]{report}
\input{metadatos}\input{plantilla/estilo}\decimalpoint
\newcommand{\un}[1]{\,\mathrm{#1}}
\begin{document}\hypersetup{pageanchor=false}\pagenumbering{gobble}
\input{plantilla/caratula}\clearpage
\hypersetup{pageanchor=true}\pagenumbering{arabic}
\begingroup\footnotesize\setlength{\cftbeforechapskip}{2pt}\setlength{\cftbeforesecskip}{0pt}\tableofcontents\endgroup\clearpage
'''+ '\n'.join(r'\input{capitulos/'+x+'}' for x in partes)+r'\end{document}'+'\n'
    (T/'main.tex').write_text(main,encoding='utf-8')


def presupuesto_tex():
    f=O/'05_Presupuesto/Presupuesto.xlsx';w=openpyxl.load_workbook(f,data_only=True);s=w['Presupuesto']
    assert isinstance(s['F28'].value,(int,float)),'Primero recalcular los libros con Excel'
    rows=[[str(s.cell(i,1).value),tex(s.cell(i,2).value),s.cell(i,3).value,n(s.cell(i,4).value),n(s.cell(i,5).value,2),n(s.cell(i,6).value,2)] for i in range(7,21)]
    for i in [22,24,25,26,28]:rows.append(['',tex(s.cell(i,2).value),'','','',n(s.cell(i,6).value,2)])
    total=s['F28'].value;w.close()
    return r'''\chapter{PRESUPUESTO REFERENCIAL}
\section{Base económica}
El presupuesto se expresa en soles, con precios estimados de octubre de 2026. No son ofertas comerciales ni cotizaciones locales. Se utilizan las cantidades del metrado desarrollado y precios por conjunto, identificando el alcance para evitar dobles conteos.

El poste instalado incluye izado y cimentación tipo dentro de la estimación económica. Los precios de 15 y 17 m se escalan por longitud respecto de la referencia de 13 m. Armados y retenidas se consideran por conjunto de montaje. Transporte y logística se incorporan mediante un factor global, sin volver a valorar el mismo concepto en las partidas unitarias.
\section{Desarrollo de costos}
\begin{align}
C_i&=Q_i\,PU_i,\qquad C_D=\sum_i C_i,\\
C_L&=0.12C_D,\qquad C_I=0.08C_D,\\
C_C&=0.20(C_D+C_L+C_I),\\
C_{\mathrm{total}}&=C_D+C_L+C_I+C_C.
\end{align}
'''+r'\Needspace{24\baselineskip}'+tabla(['Ítem','Partida','Un.','Cantidad','PU S/','Parcial S/'],rows,[.8,2.5,.5,1,1.1,1.4],'footnotesize')+r'''
\begin{equation}
\boxed{C_{\mathrm{total}}=\text{S/ }'''+n(total,2)+r'''}.
\end{equation}
El total es referencial académico, con IGV sin desagregación. No se utiliza para justificar rentabilidad ni comprometer una contratación. La contingencia es una reserva de estimación, distinta de la reserva física del conductor.
'''


def integrar():
    cuerpo=T/'main.pdf';planos=B/'planos.pdf'
    assert cuerpo.exists() and planos.exists()
    w=PdfWriter();indice=[]
    def agregar(titulo,path,pages=None):
        start=len(w.pages)+1;w.append(PdfReader(path),pages=pages,outline_item=titulo,import_outline=pages is None)
        indice.append(dict(seccion=titulo,desde=start,hasta=len(w.pages)))
    agregar('Memoria, cálculos, metrados y presupuesto',cuerpo)
    agregar('UB-01: ubicación Perú–Puno–Sandia–Limbani',B/'ubicacion/Plano_Ubicacion.pdf')
    agregar('Planos de línea primaria',planos)
    source=P/'entregables/linea-primaria-r02-presentacion/09-armados-referencia/Referencia_documental_DGE.pdf'
    # Sólo se actualiza la identidad del rótulo de proyecto, fuera del dibujo normativo.
    original=pymupdf.open(source);referencias=pymupdf.open()
    for i in range(4,27):
        referencias.insert_pdf(original,from_page=i,to_page=i)
        page=referencias[-1]
        for block in page.get_text('blocks'):
            if block[3]<24 and 'LP-R01' in block[4]:page.add_redact_annot(pymupdf.Rect(block[:4]),fill=(1,1,1))
        page.apply_redactions(images=0,graphics=0)
        page.insert_text((10,12),'LIMBANI - HUANCASAYANI | REFERENCIA DOCUMENTAL DGE',fontsize=6,fontname='helv')
    ref=B/'referencias.pdf';referencias.save(ref);referencias.close();original.close()
    agregar('Armados DGE: referencias documentales',ref)
    w.add_metadata({'/Title':'Línea primaria Limbani–Huancasayani','/Author':'Aquiles Taylor Ramos Yapo','/Subject':'Instalaciones Eléctricas II. Proyecto académico de diseño.'})
    with (O/'Expediente.pdf').open('wb') as f:w.write(f)
    (B/'indice-expediente.json').write_text(json.dumps(indice,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print('Expediente único:',len(w.pages),'páginas',flush=True)


if __name__=='__main__':
    modo=sys.argv[1] if len(sys.argv)>1 else 'memoria'
    if modo=='memoria':preparar();contenido();print(compilar(T),flush=True)
    elif modo=='integrar':integrar()
