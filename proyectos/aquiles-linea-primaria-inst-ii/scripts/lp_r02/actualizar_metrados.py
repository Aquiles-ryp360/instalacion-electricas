"""Construcción portable de M01: sólo trabaja en build hasta verificar y publicar."""
from __future__ import annotations
import argparse
import csv
import hashlib
import importlib.util
import json
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

import openpyxl
import pymupdf
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

from metrados_componentes import cargar, medir
from metrados_excel import preparar_libros
from rutas_entrega import CARPETAS, copiar_base, ruta_activa

P=Path(__file__).resolve().parents[2];ROOT=P.parents[1]
ZIP='Proyecto_Linea_Primaria_Limbani_Huancasayani.zip'
INFORME=P/'documentacion/metrados-m01/validacion-m01.json'
FUENTES_IGN=ROOT/'recursos/linea-primaria-limbani/2026-10-19/ubicacion-ign'


def insumos():
    rutas=['datos/metrados-m01/catalogo.json','datos/metrados-m01/criterios.json',
           'datos/lp-r02/parametros.json','revisiones/lp-r02/03-calculos/resultados.json']
    registro={ruta:sha(P/ruta) for ruta in rutas}
    registro.update({str(f.relative_to(ROOT)):sha(f) for f in FUENTES_IGN.glob('*') if f.is_file()})
    for ruta in ['referencias/normativa/pdf/rd024-2003-EM.pdf','recursos/linea-primaria-limbani/2026-10-19/electropuno-pac-2024.pdf']:
        registro[ruta]=sha(ROOT/ruta)
    return registro


def base():
    return json.loads((P/'revisiones/lp-r02/03-calculos/resultados.json').read_text(encoding='utf-8'))


def sha(f):return hashlib.sha256(Path(f).read_bytes()).hexdigest()


def office():
    ejecutable=shutil.which('libreoffice') or shutil.which('soffice')
    if ejecutable:return ejecutable
    for f in [Path('C:/Program Files/LibreOffice/program/soffice.exe'),Path('C:/Program Files (x86)/LibreOffice/program/soffice.exe')]:
        if f.is_file():return str(f)
    raise FileNotFoundError('Instalar LibreOffice o añadir libreoffice/soffice al PATH.')


def preparar(b):
    if b.exists():raise FileExistsError('Elegir una carpeta nueva: '+str(b))
    b.mkdir(parents=True);o=b/'Proyecto_Linea_Primaria'
    copiar_base(o)
    preparar_libros(P,o,base())
    doc=pymupdf.open(ruta_activa('Expediente.pdf'))
    toc=doc.get_toc()
    ini=next(row[2]-1 for row in toc if row[1]=='Planos de línea primaria')
    fin=next(row[2]-1 for row in toc if row[1]=='Armados DGE: referencias documentales')
    assert fin-ini==19
    planos=pymupdf.open();planos.insert_pdf(doc,from_page=ini,to_page=fin-1);planos.save(b/'planos.pdf')
    planos.close();doc.close()
    m=medir(base());(b/'resultado-metrado.json').write_text(json.dumps(m,ensure_ascii=False,indent=2)+'\n')
    with (b/'inventario-por-poste.csv').open('w',encoding='utf-8-sig',newline='') as f:
        writer=csv.DictWriter(f,fieldnames=list(m['inventario'][0]),lineterminator='\n');writer.writeheader();writer.writerows(m['inventario'])
    print('Libros e insumos preparados en',b,flush=True)


def recalcular_archivos(archivos,destino,perfil,log):
    destino.mkdir(parents=True,exist_ok=True)
    for archivo in archivos:
        esperado=destino/archivo.name
        if esperado.resolve()==archivo.resolve():raise ValueError('La exportación debe ir a una carpeta distinta del original.')
        if esperado.exists():esperado.unlink()  # Evitar validar una exportación anterior si Office falla.
    cmd=[office(),'-env:UserInstallation='+perfil.resolve().as_uri(),'--headless',
         '--convert-to','xlsx:Calc MS Excel 2007 XML','--outdir',str(destino)]+[str(x) for x in archivos]
    p=subprocess.run(cmd,stdout=subprocess.PIPE,stderr=subprocess.STDOUT,timeout=300)
    log.write_bytes(p.stdout)
    if p.returncode:raise RuntimeError(p.stdout.decode(errors='replace'))
    for x in archivos:
        if not (destino/x.name).is_file():raise FileNotFoundError('LibreOffice no exportó '+x.name)


def recalcular(b):
    o=b/'Proyecto_Linea_Primaria'
    archivos=[o/c/(n+'.xlsx') for n,c in [('Calculos','01_Calculos'),('Metrados','02_Metrados'),('Presupuesto','05_Presupuesto')]]
    destino=b/'recalculados'
    recalcular_archivos(archivos,destino,b/'perfil-libreoffice',b/'recalculo-libreoffice.txt')
    for f in archivos:shutil.copy2(destino/f.name,f)
    print('Tres libros recalculados con LibreOffice; originales de presentación conservados.',flush=True)


def resumen_pdf(f,titulo,rows,nota):
    fuentes=Path('/usr/share/fonts/truetype/dejavu')
    if not (fuentes/'DejaVuSans.ttf').is_file():
        fuentes=Path(importlib.util.find_spec('matplotlib').origin).parent/'mpl-data/fonts/ttf'
    if 'M01Sans' not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont('M01Sans',str(fuentes/'DejaVuSans.ttf')))
        pdfmetrics.registerFont(TTFont('M01Bold',str(fuentes/'DejaVuSans-Bold.ttf')))
    estilo=ParagraphStyle('m01',fontName='M01Sans',fontSize=9,leading=13)
    title=ParagraphStyle('titulo',parent=estilo,fontName='M01Bold',fontSize=15,leading=21,textColor=colors.HexColor('#254D73'))
    data=[[Paragraph(str(x),estilo) for x in row] for row in rows]
    t=Table(data,colWidths=[80*mm,28*mm,22*mm,45*mm],repeatRows=1)
    t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#EAF2F8')),('GRID',(0,0),(-1,-1),.35,colors.HexColor('#CBD7E3')),('VALIGN',(0,0),(-1,-1),'TOP'),('TOPPADDING',(0,0),(-1,-1),7),('BOTTOMPADDING',(0,0),(-1,-1),7)]))
    story=[Paragraph(titulo,title),Spacer(1,10),Paragraph('Limbani–Huancasayani · Aquiles Taylor Ramos Yapo · Metrado M01',estilo),Spacer(1,15),t,Spacer(1,15),Paragraph(nota,estilo)]
    SimpleDocTemplate(str(f),pagesize=letter,leftMargin=18*mm,rightMargin=18*mm,topMargin=20*mm,bottomMargin=20*mm).build(story)
    assert len(pymupdf.open(f))==1


def documentos(b):
    o=b/'Proyecto_Linea_Primaria';m=medir(base());z=m['conductor'];q=m['resumen_compra']
    rows=[['Concepto','Cantidad','Unidad','Desarrollo'],
          ['Postes CAC',111,'ud','72 de 13 m + 30 de 15 m + 9 de 17 m'],
          ['Armados por poste',111,'jgo','30 S + 41 A + 38 R + 2 T'],
          ['Ramales de retenida',160,'ud','Conteo mecánico independiente'],
          ['Aisladores tipo PIN',int(q['PIN']),'ud','30×3 + 41×6 + 38×1'],
          ['Discos de suspensión',int(q['DISCO']),'ud','38×12 + 2×6'],
          ['Tendido de las tres fases',f'{z["tendido_terna_m"]:,.3f}','m','Catenaria de 110 vanos'],
          ['Puentes y colas',240,'m','3×(38+2)×2'],
          ['Reserva de conductor',616,'m','3% redondeado hacia arriba'],
          ['Conductor de suministro',21131,'m','Tendido + auxiliares + reserva; redondeo final']]
    resumen_pdf(o/'02_Metrados/Resumen_Metrados.pdf','METRADOS DESARROLLADOS POR POSTE',rows,
                'El expediente contiene 111 fichas individuales. En Metrados.xlsx, selecciona el poste en Ficha_poste; consulta Matriz_armados, Inventario y Resumen_materiales. Los accesorios incluidos con pernos se controlan sin compra adicional. Adaptación CAC académica declarada en las bases del metrado.')
    w=openpyxl.load_workbook(o/'05_Presupuesto/Presupuesto.xlsx',data_only=True,read_only=True);s=w['Presupuesto']
    rows=[['Concepto','Importe S/','Unidad','Base']]
    for r in [22,24,25,26,28]:rows.append([s[f'B{r}'].value,f'{s[f"F{r}"].value:,.2f}','S/','Estimación académica por conjuntos'])
    resumen_pdf(o/'05_Presupuesto/Resumen_Presupuesto.pdf','PRESUPUESTO REFERENCIAL M01',rows,
                'La cantidad de conductor se vincula al metrado M01. Los precios de armado, retenida y puesta a tierra son por conjunto e incluyen sus componentes; el despiece no se vuelve a valorizar como costo adicional. Son precios referenciales, sin cotización comercial.')
    w.close()
    import expediente_final as e
    e.B=b;e.O=o;e.T=b/'latex'
    e.preparar();e.contenido()
    from herramientas.documentos.formato_aquiles import compilar
    print('Compilando expediente con inventarios...',flush=True)
    compilar(e.T);e.integrar()
    from memoria_descriptiva import independiente
    independiente(b)
    (o/'LEEME.txt').write_text('''LÍNEA PRIMARIA LIMBANI – HUANCASAYANI
Instalaciones Eléctricas II · Aquiles Taylor Ramos Yapo · 227368
DOCUMENTO PRINCIPAL: Expediente.pdf

METRADOS M01
El capítulo de metrados desarrolla las recetas de los armados y las 111 fichas por poste.
En 02_Metrados/Metrados.xlsx abre Ficha_poste y selecciona el identificador en B4.
Inventario contiene el detalle completo; Resumen_materiales concilia la compra global.
Accesorios_incluidos informa las piezas ya suministradas con pernos y abrazaderas.
La adaptación CAC y los coeficientes académicos se explican en las bases del metrado.

01_Calculos: cálculo eléctrico y mecánico de la base LP-R02 y resumen.
03_Planos: DXF y DWG de la entrega base, con geometría conservada.
04_Cartografia: QGIS portable y Google Earth de la entrega base.
05_Presupuesto: cantidades actualizadas, precios por conjuntos y resumen.

Tres libros recalculados con LibreOffice durante esta revisión.
El PDF debe regenerarse si se cambian entradas de los libros.
''',encoding='utf-8')


def verificar(b):
    o=b/'Proyecto_Linea_Primaria';m=medir(base());report=dict(revision_metrado='LP-R02-M01',base_tecnica='LP-R02',aprobacion_ingenieria=False,libros={})
    for name,folder in [('Calculos','01_Calculos'),('Metrados','02_Metrados'),('Presupuesto','05_Presupuesto')]:
        f=o/folder/(name+'.xlsx');a=openpyxl.load_workbook(f,data_only=False,read_only=True);v=openpyxl.load_workbook(f,data_only=True,read_only=True)
        formulas=0
        for s in a:
            for row in s:
                for x in row:
                    if x.data_type=='f':
                        formulas+=1;assert '[' not in x.value,(s.title,x.coordinate,'Vínculo externo')
        for s in v:
            for row in s:
                for x in row:assert x.data_type!='e',(name,s.title,x.coordinate,x.value)
        assert not a._external_links
        if name in ('Metrados','Presupuesto'):
            for cel,valor in [('E10',111),('E17',111),('E19',160),('E25',240),('E26',616),('E27',21131)]:assert abs(v['Resumen'][cel].value-valor)<1e-6,(name,cel,v['Resumen'][cel].value)
        if name=='Calculos':
            f0=next(x for x in base()['comparacion_tensiones'] if x['kV']==10 and x['toma_pu']==1.05)
            assert abs(v['Resumen']['D16'].value-f0['BT_llegada_V'])<1e-6
            assert abs(v['Resumen']['D18'].value-f0['P_generacion_kW'])<1e-6
        if name=='Metrados':
            resumen={row[0]:row[3] for row in v['Resumen_materiales'].iter_rows(min_row=8,values_only=True) if row[0]}
            for k,q in m['resumen_compra'].items():assert abs(resumen[k]-q)<1e-6,(k,resumen[k],q)
            inc={row[0]:row[3] for row in v['Accesorios_incluidos'].iter_rows(min_row=8,values_only=True) if row[0]}
            for k,q in m['resumen_incluidos'].items():assert abs(inc[k]-q)<1e-6,(k,inc[k],q)
            inv={}
            for row in v['Inventario'].iter_rows(min_row=8,values_only=True):
                if row[0]:inv[(row[0],row[1])]=row[9]
            from collections import Counter
            esperado_por_clave=Counter()
            for e in m['inventario']:esperado_por_clave[(e['poste'],e['material'])]+=e['cantidad']
            for poste in base()['postes']:
                for k in cargar()[0]['materiales']:
                    esperado=esperado_por_clave[(poste['id'],k)]
                    assert abs(inv[(poste['id'],k)]-esperado)<1e-6,(poste['id'],k,inv[(poste['id'],k)],esperado)
        if name=='Presupuesto':
            s=v['Presupuesto'];report['presupuesto_referencial_pen']=s['F28'].value
            for r in range(7,21):assert abs(s[f'F{r}'].value-round(s[f'D{r}'].value*s[f'E{r}'].value+1e-8,2))<1e-6
        report['libros'][name]=dict(hojas=a.sheetnames,formulas=formulas,sha256=sha(f),motor_recalculo='LibreOffice',errores=0)
        a.close();v.close()
    pdf=pymupdf.open(o/'Expediente.pdf');text='\n'.join(p.get_text() for p in pdf)
    for p in base()['postes']:assert 'FICHA DE METRADO' in text and p['id'] in text
    fichas=[row for row in pdf.get_toc() if 'Inventario' in row[1] and row[1].startswith('P')]
    assert len(fichas)==111,len(fichas)
    for i,ficha in enumerate(fichas):
        assert ficha[1].startswith(m['asignaciones'][i]['poste']+' ')
        assert 'FICHA DE METRADO' in pdf[ficha[2]-1].get_text()
        if i<110:assert fichas[i+1][2]-ficha[2]==1,('Ficha dividida',ficha[1])
    siguiente=min(row[2] for row in pdf.get_toc() if row[2]>fichas[-1][2])
    assert siguiente==fichas[-1][2]+1,('Última ficha dividida',fichas[-1][1])
    assert '21 873' not in text and '21,873' not in text
    for folder,name in [('01_Calculos','Resumen_Calculos'),('02_Metrados','Resumen_Metrados'),('05_Presupuesto','Resumen_Presupuesto')]:assert len(pymupdf.open(o/folder/(name+'.pdf')))==1
    # CAD y cartografía nativos se preservan; no se atribuye una nueva reapertura nativa.
    preservados=[]
    for folder in ['03_Planos','04_Cartografia']:
        for f in (o/folder).rglob('*'):
            if f.is_file():
                original=ruta_activa(f.relative_to(o))
                assert sha(f)==sha(original),f
                preservados.append(str(f.relative_to(o)))
    archivos=sorted(f for f in o.rglob('*') if f.is_file() and not f.name.startswith('.'))
    assert len(archivos)==16,len(archivos)
    report.update(paginas_expediente=len(pdf),fichas_individuales=111,ficha_por_pagina=True,archivos_nativos_sin_cambios=preservados,
                  conductor=m['conductor'],materiales_compra=m['resumen_compra'],accesorios_incluidos=m['resumen_incluidos'],
                  giros_especiales=m['giros_especiales'],motor_version=subprocess.check_output([office(),'--version'],text=True).strip(),
                  insumos_sha256=insumos(),entregables_sha256={str(f.relative_to(o)):sha(f) for f in archivos})
    pdf.close()
    (b/'validacion-m01.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    print('Verificación M01 aprobada:',report['paginas_expediente'],'páginas, 111 fichas y conciliación de todos los componentes.',flush=True)
    return report


def organizar(b):
    """Carpetas de lectura solicitadas, sin duplicar el expediente en la raíz."""
    o=b/'Proyecto_Linea_Primaria';d=b/'Entrega_M01'
    if d.exists():shutil.rmtree(d)  # Sólo salida regenerable dentro de build.
    d.mkdir();(d/'00_Expediente').mkdir();(d/'01_Memoria_Descriptiva').mkdir()
    shutil.copy2(o/'Expediente.pdf',d/'00_Expediente/Expediente.pdf')
    shutil.copy2(b/'Memoria_Descriptiva.pdf',d/'01_Memoria_Descriptiva/Memoria_Descriptiva.pdf')
    for vieja,nueva in CARPETAS.items():shutil.copytree(o/vieja,d/nueva)
    for ext in ['pdf','svg']:shutil.copy2(b/'ubicacion'/('Plano_Ubicacion.'+ext),d/'04_Planos'/('Plano_Ubicacion.'+ext))
    recursos=d/'07_Recursos';recursos.mkdir()
    shutil.copytree(FUENTES_IGN,recursos/'Cartografia_IGN')
    shutil.copytree(P/'datos/metrados-m01',recursos/'Bases_Metrados')
    (recursos/'Normativa').mkdir()
    shutil.copy2(ROOT/'referencias/normativa/pdf/rd024-2003-EM.pdf',recursos/'Normativa/RD024_2003_EM.pdf')
    shutil.copy2(ROOT/'recursos/linea-primaria-limbani/2026-10-19/electropuno-pac-2024.pdf',recursos/'Normativa/ElectroPuno_PAC2024.pdf')
    cat,_=cargar();lineas=['FUENTES Y ALCANCE DE USO · M01','',
        'IGN: límites referenciales oficiales; plano UB-01 elaborado por el proyecto.',
        'https://www.idep.gob.pe/geoportal/rest/services/DATOS_GEOESPACIALES/L%C3%8DMITES/MapServer',
        'Consulta 19/10/2026. Cartografia_IGN conserva respuestas y procedencia con SHA256.',
        'Se excluye del dibujo la entidad departamental vacía sin nombre/código; el original se conserva.',
        'El eje y las posiciones de equipo proceden de LP-R02; no de un levantamiento del IGN.','']
    for fuente in cat['fuentes']:lineas += [fuente['titulo'],fuente['url'],'SHA256: '+fuente['sha256'],'Uso: '+fuente['uso'],'']
    lineas+=['Bases_Metrados conserva coeficientes, unidades, hipótesis y compatibilidades pendientes.',
             'Los antecedentes censales INEI y restantes fuentes de cálculo se citan en el expediente y en el repositorio completo.',
             'El recálculo eléctrico/mecánico sigue la base LP-R02. Las masas reales del despiece no están certificadas.']
    (recursos/'LEEME_Fuentes.txt').write_text('\n'.join(lineas)+'\n',encoding='utf-8')
    (d/'LEEME.txt').write_text('''LÍNEA PRIMARIA LIMBANI–HUANCASAYANI · M01
Instalaciones Eléctricas II · Aquiles Taylor Ramos Yapo · 227368

00_Expediente: expediente integrado, con 111 fichas de inventario por poste.
01_Memoria_Descriptiva: memoria independiente, con ubicación y lámina UB-01 A3.
02_Calculos: libro eléctrico/mecánico formulado y resumen PDF.
03_Metrados: libro de 12 hojas y resumen PDF. Selecciona el poste en Ficha_poste!B4.
04_Planos: CAD base LP-R02 y plano UB-01 de ubicación en PDF/SVG editable.
05_Cartografia: QGIS portable, DEM, GeoPackage y Google Earth de la base.
06_Presupuesto: cantidades M01 vinculadas y precios referenciales por conjunto.
07_Recursos: fuentes oficiales, cartografía IGN y bases del metrado.

En Metrados.xlsx: Matriz_armados → Inventario → Ficha_poste → Resumen_materiales.
Accesorios_incluidos informa piezas ya suministradas con pernos/abrazaderas.
Los tres libros se recalcularon con LibreOffice. Si editas entradas, regenera el PDF.
Cada XLSX es autónomo: editar uno no modifica los otros. Regenera los tres desde
las entradas canónicas para sincronizar sus cantidades; no hay vínculos externos.
Las adaptaciones CAC son académicas y se explican en las bases del metrado.
El plano usa límites referenciales oficiales; el corredor es de gabinete.
''',encoding='utf-8')
    return d


def empaquetar(b):
    report=verificar(b);o=organizar(b)
    archivos=sorted(f for f in o.rglob('*') if f.is_file() and not f.name.startswith('.'))
    assert len([f for f in o.iterdir() if f.is_dir()])==8
    z=b/ZIP
    with zipfile.ZipFile(z,'w',zipfile.ZIP_DEFLATED,compresslevel=6) as out:
        for f in archivos:out.write(f,Path('Proyecto_Linea_Primaria')/f.relative_to(o))
    with zipfile.ZipFile(z) as out:assert out.testzip() is None
    with pymupdf.open(o/'01_Memoria_Descriptiva/Memoria_Descriptiva.pdf') as memoria:
        assert any('UB-01' in r[1] for r in memoria.get_toc())
        paginas_memoria=len(memoria)
    report.update(zip_sha256=sha(z),archivos_zip=len(archivos),carpetas_entrega=8,paginas_memoria=paginas_memoria,
                  ubicacion=json.loads((b/'ubicacion/validacion-ubicacion.json').read_text()),
                  entregables_sha256={str(f.relative_to(o)):sha(f) for f in archivos})
    (b/'validacion-m01.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print('ZIP verificado:',z,sha(z),flush=True)


def verificar_publicado(_b=None):
    """Verificación portable de la entrega revisada, sin necesitar build ni Office."""
    report=json.loads(INFORME.read_text(encoding='utf-8'))
    assert report['insumos_sha256']==insumos(),'Las entradas cambiaron: regenerar y revisar la entrega.'
    o=P/'entregables/proyecto-final'
    presentes={str(f.relative_to(o)) for f in o.rglob('*') if f.is_file() and not f.name.startswith('.')}
    assert presentes==set(report['entregables_sha256']), 'Cambió el conjunto de archivos de la entrega.'
    for nombre,h in report['entregables_sha256'].items():assert sha(o/nombre)==h,nombre
    z=P/'entregables'/ZIP
    assert sha(z)==report['zip_sha256'],'El ZIP no coincide con la revisión.'
    with zipfile.ZipFile(z) as paquete:
        assert paquete.testzip() is None
        esperados={'Proyecto_Linea_Primaria/'+n:h for n,h in report['entregables_sha256'].items()}
        assert set(paquete.namelist())==set(esperados)
        for nombre,h in esperados.items():assert hashlib.sha256(paquete.read(nombre)).hexdigest()==h,nombre
    with pymupdf.open(o/'00_Expediente/Expediente.pdf') as pdf:
        assert len(pdf)==report['paginas_expediente']
        fichas=[r for r in pdf.get_toc() if r[1].startswith('P') and 'Inventario' in r[1]]
        assert len(fichas)==111
    print('Entrega M01 íntegra:',report['archivos_zip'],'archivos en 8 carpetas; 111 fichas, insumos y ZIP conciliados.',flush=True)


def ensayos(b):
    """Ediciones de prueba en una copia; nunca modifica los libros publicados."""
    origen=b/'Proyecto_Linea_Primaria/02_Metrados/Metrados.xlsx'
    inicial=sha(origen);prueba=b/'ensayo-entradas';prueba.mkdir(exist_ok=True)
    f=prueba/'Ensayo_Metrados.xlsx';w=openpyxl.load_workbook(origen)
    w['Postes']['I8']='ANCLAJE';w['Postes']['E8']=15;w['Ficha_poste']['B4']='P002'
    for row in w['Matriz_armados'].iter_rows(min_row=8):
        if row[0].value=='DISCO':w['Matriz_armados'].cell(row[0].row,6,13)
    w.save(f);w.close()
    out=prueba/'recalculado'
    recalcular_archivos([f],out,prueba/'perfil',prueba/'libreoffice.txt')
    v=openpyxl.load_workbook(out/f.name,data_only=True,read_only=True)
    s=v['Resumen'];esperados={'E7':71,'E8':31,'E9':9,'E10':111,'E19':162,'E25':246,'E27':21137}
    for cell,q in esperados.items():assert abs(s[cell].value-q)<1e-6,(cell,s[cell].value,q)
    total={row[0]:row[3] for row in v['Resumen_materiales'].iter_rows(min_row=8,values_only=True) if row[0]}
    assert total['PIN']==372 and total['DISCO']==519 and total['CRU-M']==193,total
    ficha={row[0]:row[7] for row in v['Ficha_poste'].iter_rows(min_row=8,values_only=True) if row[0]}
    assert ficha['DISCO']==13 and ficha['PIN']==1 and ficha['BLO-R']==2,ficha
    assert ficha['P13']==0 and ficha['P15']==1 and ficha['AAAC-P']==6
    v.close();assert sha(origen)==inicial
    report=dict(motor='LibreOffice',ediciones=['P002 suspensión -> anclaje','P002 altura 13 -> 15 m','DISCO por CAC-R 12 -> 13','Ficha_poste selecciona P002'],
                resultados_esperados=esperados,materiales_comprobados={'PIN':372,'DISCO':519,'CRU-M':193},
                ficha_P002={'DISCO':13,'PIN':1,'BLO-R':2,'P15':1},libro_entrega_sin_modificaciones=True)
    (b/'ensayos-entradas.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    print('Ensayos de entradas y selector de ficha aprobados; copia de entrega intacta.',flush=True)


if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('etapa',choices=['preparar','recalcular','documentos','verificar','empaquetar','ensayos','verificar-publicado'])
    ap.add_argument('--trabajo',type=Path,default=ROOT/'build'/P.name/'metrados-m01');a=ap.parse_args()
    globals()[a.etapa.replace('-','_')](a.trabajo.resolve())
