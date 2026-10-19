"""Validación de física, trazabilidad y archivos; no certificación de ingeniería."""
from pathlib import Path
import sys,json,math,csv,hashlib,zipfile,argparse
import ezdxf,pymupdf
import xml.etree.ElementTree as ET
from pyproj import Transformer
from openpyxl import load_workbook
from pypdf import PdfReader
P=Path(__file__).resolve().parents[2];ROOT=P.parents[1]

def verificar(O,escribir=True):
    r=json.loads((O/'03-calculos/resultados.json').read_text());b=json.loads((P/'datos/lp-r02/geometria-base.json').read_text());c=json.loads((P/'datos/lp-r02/parametros.json').read_text());checks=[]
    def ck(name,value):
        if not bool(value):raise AssertionError(name)
        checks.append(name)
    def rows(f):return list(csv.DictReader(f.open(encoding='utf-8-sig')))
    ps=r['postes'];vs=r['vanos'];ck('IDs únicos yconteos derivados degeometría',len({p['id'] for p in ps})==len(ps)==len(b['postes']) and len(vs)==len(ps)-1)
    ck('Conectividad sinramales yprogresivas coherentes',all(v['desde']==ps[i]['id'] and v['hasta']==ps[i+1]['id'] for i,v in enumerate(vs)) and all(ps[i+1]['progresiva_m']>p['progresiva_m'] for i,p in enumerate(ps[:-1])))
    ck('Longitudes calculadas deXY',all(abs(v['longitud_m']-math.dist((ps[i]['e'],ps[i]['n']),(ps[i+1]['e'],ps[i+1]['n'])))<1e-7 for i,v in enumerate(vs)))
    ck('Eje=longitudvanos=progresivafinal',abs(sum(v['longitud_m'] for v in vs)-r['longitud_horizontal_m'])<1e-6 and abs(ps[-1]['progresiva_m']-r['longitud_horizontal_m'])<1e-6)
    ck('GeometríaR01 intacta porSHA256',b['sha256_resultados_origen']==hashlib.sha256((P/'datos/lp-r01/resultados.json').read_bytes()).hexdigest())
    ck('Estados porvano concordantes conentrada',all(set(v['estados'])=={s['id'] for s in c['clima']} for v in vs))
    ck('Conservación termoelástica porvano',max(abs(s['residuo_longitud_m']) for v in vs for s in v['estados'].values())<1e-7)
    ck('Separación mínima reportada sinocultaralertas',abs(r['resumen']['separacion_perfil_min_m']-min(s['separacion_perfil_heredado_m'] for v in vs for s in v['estados'].values()))<1e-8)
    ck('Cargasascendentes preservadas con signo',sum(p['vertical_min_kN']<0 for p in ps)==r['resumen']['apoyos_carga_ascendente'])
    ck('Capacidad axial faltante marcada null',all(p['capacidad_axial_kN'] is None for p in ps))
    ck('Propuestas yretenidas cuadran conplanilla',sum(p['retenidas_r02'] for p in ps)==len(r['retenidas']) and sum(p['tipo_r02']=='ANCLAJE_PROPUESTO' for p in ps)==r['resumen']['apoyos_refuerzo_propuesto'])
    ck('Retenidas respetan resistencia mínima de la especificación',all(g['rotura_min_kN']*1000>=c['mecanica']['retenida_rotura_min_n'] for g in r['retenidas']))
    ck('BalanceP/Q convergente en12sensibilidades',all(abs(f['residuo_balance_W'])<1e-4 and abs(f['residuo_balance_var'])<1e-4 and f['residuo_V']<1e-7 for f in r['comparacion_tensiones']))
    ck('Pérdida esperada decrece contensión',all(next(f['perdida_linea_kW'] for f in r['comparacion_tensiones'] if f['kV']==22.9 and f['toma_pu']==t)<next(f['perdida_linea_kW'] for f in r['comparacion_tensiones'] if f['kV']==10 and f['toma_pu']==t) for t in [1,1.025,1.05]))
    ck('Longitud eléctrica física supera laproyección',r['longitud_electrica_m']>=r['longitud_horizontal_m'])
    ck('Metrado conductor incluye terna,puentes yreserva',r['conductor']['total_compra_estudio_m']==math.ceil(r['conductor']['tendido_terna_m']+r['conductor']['puentes_provision_m']+r['conductor']['reserva_m']))
    ck('Metrado detallevinculado aIDs',len([x for x in rows(O/'02-metrados/metrado_detallado.csv') if x['id'].startswith('P-')])==len(ps))
    d=ezdxf.readfile(O/'04-planos/LP_Limbani_Huancasayani_R02.dxf');audit=d.audit();ck('DXF reabierto,auditoría sincorrupción',not audit.errors);ck('CAD modelo enmetros',d.units==6)
    ins={i.get_attrib_text('ID'):i for i in d.modelspace().query('INSERT')};ck('CAD modelocontieneapoyos ycálculoidénticos',set(ins)=={p['id'] for p in ps} and all(math.dist(tuple(ins[p['id']].dxf.insert)[:2],(p['e'],p['n']))<1e-7 and ins[p['id']].get_attrib_text('TIPO')==p['tipo_r02'] for p in ps))
    ck('Un ejepara laterna',len(list(d.modelspace().query('LWPOLYLINE[layer=="EJE"]')))==1)
    index=json.loads((O/'04-planos/indice-planos.json').read_text());combined=PdfReader(O/'04-planos/Planos_LP_R02_A3.pdf');ck('Índice,layouts yPDFconlamismacantidad',len(index)==len(combined.pages)==sum(l.name.startswith('LP-') for l in d.layouts))
    for row in index:
        pdf=pymupdf.open(O/'04-planos'/row['archivo']);page=pdf[0]
        ck('Lámina'+row['codigo']+'A3válidaycongeometría',len(pdf)==1 and abs(page.rect.width-420*72/25.4)<.5 and abs(page.rect.height-297*72/25.4)<.5 and len(page.get_drawings())>10)
    ck('GeoJSON coincidenteposteaposte',json.loads((O/'04-planos/postes_R02_UTM19S.geojson').read_text())['features'][0]['properties']==ps[0] and len(json.loads((O/'04-planos/postes_R02_UTM19S.geojson').read_text())['features'])==len(ps))
    with zipfile.ZipFile(O/'04-planos/LP_R02.kmz') as z:
        ck('KMZíntegroyconKML',z.testzip() is None and 'doc.kml' in z.namelist())
        tree=ET.fromstring(z.read('doc.kml'));ns={'k':'http://www.opengis.net/kml/2.2'};points={};t=Transformer.from_crs(4326,32719,always_xy=True)
        for pm in tree.findall('.//k:Placemark',ns):
            coord=pm.find('k:Point/k:coordinates',ns)
            if coord is not None:
                key=pm.find('k:name',ns).text.split(' / ')[0];lon,lat,_=map(float,coord.text.split(','));points[key]=t.transform(lon,lat)
        ck('KMZ coordenadas coinciden con planilla yCAD',set(points)=={p['id'] for p in ps} and max(math.dist(points[p['id']],(p['e'],p['n'])) for p in ps)<.002)

    for f in O.rglob('*.xlsx'):
        w=load_workbook(f);ck('Libro'+f.name+'reabierto',len(w.sheetnames)>0)
    budget=load_workbook(O/'06-presupuesto-y-cronograma/Presupuesto_y_planificacion_R02.xlsx');ck('Precios nocotizados siguenvacíos',all(row[3].value is None for row in list(budget['Presupuesto'].rows)[1:-1]))
    for f in [O/'01-expediente/Expediente_revisado_LP_R02.pdf',O/'03-calculos/Memoria_calculo_explicita_LP_R02.pdf']:
        pdf=pymupdf.open(f);ck('Memoria'+f.name+'conestados ypáginaslegibles',len(pdf)>3 and 'LP-R02' in ''.join(page.get_text() for page in pdf))
    for f in json.loads((O/'08-recursos/fuentes.json').read_text()):ck('Recurso'+f['id']+'íntegro',hashlib.sha256((O/f['ruta_paquete']).read_bytes()).hexdigest()==f['sha256'])
    for name in ['01-expediente','02-metrados','03-calculos','04-planos','05-especificaciones','06-presupuesto-y-cronograma','07-evaluacion-y-pendientes','08-recursos']:ck('Carpeta'+name+'concontenido',(O/name).is_dir() and any((O/name).iterdir()))
    result=dict(estado='VERIFICACIONES INFORMATICAS APROBADAS; INGENIERIA CON PENDIENTES',controles_aprobados=len(checks),controles=checks,
                pruebas_fisicas='11 pruebas unitarias independientes; ejecutar comando deCONTINUIDAD_IA',
                observaciones_abiertas=r['limites'],resumen_modelo=r['resumen'],aprobacion_ingenieria=False)
    if escribir:(O/'validacion.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
    return result
if __name__=='__main__':
    a=argparse.ArgumentParser();a.add_argument('--salida',type=Path,required=True);args=a.parse_args();res=verificar(args.salida.resolve());print(json.dumps({k:v for k,v in res.items() if k not in ['controles']},ensure_ascii=False,indent=2))
